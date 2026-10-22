import io
import streamlit as st
from pdf2docx import Converter
from docx import Document
from docx2pdf import convert
import os
import tempfile

st.title("📄 PDF Editor & Converter")
st.write("Ubah PDF ke Word, edit teksnya secara langsung, lalu simpan kembali sebagai PDF!")

# Step 1: Upload PDF
uploaded_file = st.file_uploader("Upload file PDF Anda", type=["pdf"])

if uploaded_file is not None:
    # Simpan PDF ke file sementara
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_pdf:
        tmp_pdf.write(uploaded_file.read())
        tmp_pdf_path = tmp_pdf.name

    tmp_docx_path = tmp_pdf_path.replace(".pdf", ".docx")

    # Konversi PDF ke DOCX menggunakan pdf2docx
    with st.spinner("Mengonversi PDF ke DOCX..."):
        cv = Converter(tmp_pdf_path)
        cv.convert(tmp_docx_path, start=0, end=None)
        cv.close()

    # Baca isi dokumen DOCX
    doc = Document(tmp_docx_path)
    full_text = "\n".join([para.text for para in doc.paragraphs])

    st.subheader("📝 Edit Teks Dokumen")
    # Area untuk mengedit teks yang diekstrak dari PDF
    edited_text = st.text_area("Edit konten di bawah ini:", value=full_text, height=350)

    # Step 2: Simpan perubahan dan konversi balik ke PDF
    if st.button("🔄 Simpan & Konversi ke PDF Baru"):
        with st.spinner("Memproses PDF baru..."):
            # Buat dokumen DOCX baru dari teks yang disunting
            new_doc = Document()
            for line in edited_text.split("\n"):
                new_doc.add_paragraph(line)

            output_docx = "edited_document.docx"
            output_pdf = "edited_document.pdf"
            new_doc.save(output_docx)

            # Konversi DOCX kembali ke PDF
            try:
                # Opsi A: Menggunakan docx2pdf (Membutuhkan MS Word terinstall pada Windows/Mac)
                convert(output_docx, output_pdf)
                
                with open(output_pdf, "rb") as pdf_file:
                    st.success("Berhasil dikonversi kembali ke PDF!")
                    st.download_button(
                        label="📥 Download PDF Hasil Edit",
                        data=pdf_file,
                        file_name="hasil_edit.pdf",
                        mime="application/pdf"
                    )
            except Exception as e:
                st.error(f"Gagal mengonversi ke PDF: {e}")
                st.info("Catatan: `docx2pdf` memerlukan Microsoft Word terinstall di komputer server/lokal.")

    # Cleanup file sementara
    if os.path.exists(tmp_pdf_path):
        os.remove(tmp_pdf_path)
    if os.path.exists(tmp_docx_path):
        os.remove(tmp_docx_path)
